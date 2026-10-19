import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

BASE_DIR = r"D:\Project\ruangtumbuh"
DOCS_DIR = os.path.join(BASE_DIR, "docs")
SCREENSHOTS_DIR = os.path.join(DOCS_DIR, "screenshots")
os.makedirs(DOCS_DIR, exist_ok=True)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'))

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def setup_document(role_title):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hrun = hp.add_run(f"Ruang Tumbuh | {role_title} • ruangtumbuh.appocket.id")
    hrun.font.size = Pt(8.5)
    hrun.font.color.rgb = RGBColor(148, 163, 184)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    frun = fp.add_run("Buku Panduan Pengguna Resmi - Platform Ruang Tumbuh © 2026")
    frun.font.size = Pt(8.5)
    frun.font.color.rgb = RGBColor(148, 163, 184)
    return doc

def add_callout(doc, text, title="💡 CATATAN PENTING / TIPS", border_color="0284C7", bg_color="F0F9FF"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Inches(6.4)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run_title = p.add_run(f"{title}\n")
    run_title.bold = True
    run_title.font.size = Pt(9.5)
    run_title.font.color.rgb = RGBColor(2, 132, 199) if border_color == "0284C7" else RGBColor(185, 28, 28) if border_color in ["E11D48", "DC2626"] else RGBColor(13, 148, 136)
    run_text = p.add_run(text)
    run_text.font.size = Pt(9)
    run_text.font.color.rgb = RGBColor(51, 65, 85)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_table_data(doc, headers, data, col_widths, theme_color="1E3A8A"):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        hdr_cells[i].width = col_widths[i]
        set_cell_background(hdr_cells[i], theme_color)
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=120, right=120)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            r.font.size = Pt(9)
            
    for row_idx, row_data in enumerate(data):
        row_cells = table.add_row().cells
        bg = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for i, val in enumerate(row_data):
            row_cells[i].text = str(val)
            row_cells[i].width = col_widths[i]
            set_cell_background(row_cells[i], bg)
            set_cell_margins(row_cells[i], top=90, bottom=90, left=120, right=120)
            p = row_cells[i].paragraphs[0]
            for r in p.runs:
                r.font.size = Pt(8.5)
                r.font.color.rgb = RGBColor(30, 41, 59)
                
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def add_screenshot(doc, img_name, caption):
    img_path = os.path.join(SCREENSHOTS_DIR, img_name)
    if os.path.exists(img_path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run()
        run.add_picture(img_path, width=Inches(5.8))
        
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_before = Pt(2)
        cp.paragraph_format.space_after = Pt(8)
        crun = cp.add_run(caption)
        crun.italic = True
        crun.font.size = Pt(8.5)
        crun.font.color.rgb = RGBColor(100, 116, 139)

def add_cover(doc, title, subtitle, module_name):
    p_pre = doc.add_paragraph()
    p_pre.paragraph_format.space_before = Pt(24)
    p_pre.paragraph_format.space_after = Pt(4)
    r_pre = p_pre.add_run("PANDUAN OPERASIONAL & PENGGUNAAN SISTEM")
    r_pre.font.size = Pt(10.5)
    r_pre.bold = True
    r_pre.font.color.rgb = RGBColor(14, 165, 233)

    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(6)
    r_title = p_title.add_run(title)
    r_title.bold = True
    r_title.font.size = Pt(24)
    r_title.font.color.rgb = RGBColor(15, 23, 42)

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(14)
    r_sub = p_sub.add_run(subtitle)
    r_sub.font.size = Pt(12)
    r_sub.font.color.rgb = RGBColor(71, 85, 105)

    headers = ["Parameter", "Keterangan Informasi"]
    data = [
        ["Nama Aplikasi", "Ruang Tumbuh"],
        ["Tautan Aplikasi (URL)", "https://ruangtumbuh.appocket.id"],
        ["Modul yang Dibahas", module_name],
        ["Tanggal Pembaruan Terakhir", "Oktober 2026"],
        ["Format Distribusi Dokumen", "Microsoft Word (.docx) & PDF Ready"],
    ]
    add_table_data(doc, headers, data, [Inches(2.2), Inches(4.2)], theme_color="0F172A")
    doc.add_page_break()

def add_heading_1(doc, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(16)
    h.paragraph_format.space_after = Pt(6)
    r = h.add_run(text)
    r.bold = True
    r.font.size = Pt(14.5)
    r.font.color.rgb = RGBColor(15, 23, 42)
    return h

def add_heading_2(doc, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(11)
    h.paragraph_format.space_after = Pt(4)
    r = h.add_run(text)
    r.bold = True
    r.font.size = Pt(11.5)
    r.font.color.rgb = RGBColor(30, 58, 138)
    return h

def add_heading_3(doc, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(7)
    h.paragraph_format.space_after = Pt(2)
    r = h.add_run(text)
    r.bold = True
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(51, 65, 85)
    return h

def add_paragraph(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(51, 65, 85)
    return p

# ==============================================================================
# 1. GENERATE DOKUMEN ADMIN
# ==============================================================================
def generate_admin_doc():
    print("Membangun Dokumen Baru: Panduan_Pengguna_Ruang_Tumbuh_Admin.docx...")
    doc = setup_document("Buku Panduan Administrator")
    add_cover(
        doc,
        title="Buku Panduan Pengguna: Administrator",
        subtitle="Platform Konseling & Resiliensi Remaja Berbasis Kearifan Lokal Nusantara",
        module_name="Modul Pengelolaan Master Data & Konfigurasi Sistem (Admin)"
    )

    # 1. Judul & Versi
    add_heading_1(doc, "1. Informasi Modul & Akses Sistem")
    add_paragraph(doc, "Dokumen ini merupakan panduan teknis operasional bagi Administrator sistem platform Ruang Tumbuh. Melalui modul ini, Administrator memiliki wewenang penuh dalam mengelola pangkalan data pengguna, daftar konselor bimbingan konseling, instrumen asesmen psikologis (Psychological Well-Being 25 butir & Deteksi Risiko Ide Bunuh Diri 12 butir), modul self-help berbasis kearifan budaya Nusantara, serta modul edukasi materi.")
    
    headers_info = ["Komponen", "Spesifikasi / Keterangan"]
    data_info = [
        ["Nama Aplikasi", "Ruang Tumbuh"],
        ["Domain Akses", "https://ruangtumbuh.appocket.id"],
        ["Role Sasaran", "Administrator Sistem (Admin)"],
        ["Modul Utama", "Manajemen User, Konselor, Asesmen Dua Bagian, Self-Help Budaya, & Pembelajaran"],
        ["Periode Pembaruan", "Oktober 2026"],
    ]
    add_table_data(doc, headers_info, data_info, [Inches(2.0), Inches(4.4)], theme_color="0F172A")

    # 2. Pendahuluan Ringan
    add_heading_1(doc, "2. Pendahuluan Ringan")
    add_paragraph(doc, "Ruang Tumbuh adalah platform bimbingan konseling online dan penguatan resiliensi remaja yang mengintegrasikan pendekatan klinis modern dengan nilai-nilai kearifan lokal budaya Indonesia. Sebagai Administrator, Anda bertindak sebagai pengendali utama tata kelola data sistem.")
    add_paragraph(doc, "Fokus tugas Administrator mencakup:\n"
                  "1. Pengelolaan Akun Pengguna: Memverifikasi akun siswa (konseli), mendaftarkan konselor baru, dan mengatur hak akses akun.\n"
                  "2. Profil & Urutan Konselor: Menata nomor urut konselor, deskripsi keahlian, nomor WhatsApp, serta foto profil.\n"
                  "3. Pengelolaan Instrumen Psikologis (2 Bagian Baku):\n"
                  "   • Bagian 1: Psychological Well-Being Carol D. Ryff (25 butir, rentang 25–125 poin, titik potong skrining <= 80).\n"
                  "   • Bagian 2: Identifikasi Diri / Ide Bunuh Diri (12 butir, rentang 12–60 poin, ambang krisis >= 41).\n"
                  "4. Pengayaan Kearifan Budaya: Mengelola suku-suku nusantara, materi narasi resiliensi SFBC, dan pertanyaan refleksi mandiri.")

    add_callout(doc, "Akun Administrator memiliki hak istimewa tingkat tinggi. Jaga kerahasiaan kata sandi dan hindari login dari komputer umum tanpa pengamanan.", title="⚠️ KEAMANAN HAK AKSES ADMINISTRATOR", border_color="E11D48", bg_color="FFF1F2")

    # 3. Prasyarat (Prerequisites)
    add_heading_1(doc, "3. Prasyarat (Prerequisites)")
    prereq_headers = ["Kebutuhan", "Spesifikasi Minimal yang Disarankan"]
    prereq_data = [
        ["Akun & Kredensial", "Akun terdaftar role 'admin' (email: admin@ruangtumbuh.com)."],
        ["Perangkat", "Laptop atau PC Desktop dengan layar minimal resolusi 1280 x 720 piksel."],
        ["Peramban Web", "Google Chrome, Microsoft Edge, atau Mozilla Firefox versi mutakhir."],
        ["Koneksi Internet", "Koneksi stabil minimal 2 Mbps."]
    ]
    add_table_data(doc, prereq_headers, prereq_data, [Inches(2.2), Inches(4.2)], theme_color="1E3A8A")

    # 4. Langkah Penggunaan (Step-by-Step)
    add_heading_1(doc, "4. Langkah Penggunaan (Step-by-Step)")
    
    add_heading_2(doc, "4.1 Masuk ke Halaman Login & Autentikasi Admin")
    add_paragraph(doc, "Untuk masuk ke sistem:\n"
                  "1. Buka browser dan arahkan ke: https://ruangtumbuh.appocket.id/login\n"
                  "2. Masukkan email Administrator (admin@ruangtumbuh.com) dan password.\n"
                  "3. Klik tombol 'Log in'. Sistem otomatis membuka Dashboard Administrator.")
    add_screenshot(doc, "02_login_page.png", "Gambar 1.1: Halaman Masuk (Login) Ruang Tumbuh")

    add_heading_2(doc, "4.2 Tinjauan Antarmuka Dashboard Administrator")
    add_paragraph(doc, "Dashboard Administrator (/admin) memberikan ikhtisar menyeluruh terhadap operasional sistem, navigasi cepat pada bilah sisi kiri, serta indikator aktivitas.")
    add_screenshot(doc, "admin_01_dashboard.png", "Gambar 1.2: Antarmuka Dashboard Administrator Ruang Tumbuh")

    add_heading_2(doc, "4.3 Manajemen Akun Pengguna (Users)")
    add_paragraph(doc, "Pada menu 'Users' (/admin/users), Anda dapat memantau seluruh pengguna terdaftar, menyaring berdasarkan role (Admin, Konselor, Konseli), menambahkan pengguna baru, serta melakukan pengubahan kata sandi atau data profil.")
    add_screenshot(doc, "admin_02_users.png", "Gambar 1.3: Halaman Manajemen Pengguna (Users)")

    add_heading_2(doc, "4.4 Manajemen Data & Profil Konselor")
    add_paragraph(doc, "Menu 'Counselors' (/admin/counselors) digunakan untuk mengelola konselor yang akan tampil di hadapan siswa. Atur nomor urut penampilan (Order Number), nomor kontak WhatsApp untuk darurat/konsultasi cepat, serta motto dan biografi keilmuan.")
    add_screenshot(doc, "admin_03_counselors.png", "Gambar 1.4: Halaman Manajemen Data Konselor")

    add_heading_2(doc, "4.5 Manajemen Bank Soal Instrumen (Questions)")
    add_paragraph(doc, "Menu 'Questions' (/admin/questions) memuat bank soal instrumen terstandar:\n"
                  "• Bagian 1: 25 butir Psychological Well-Being (kategori: well_being) yang mengukur 6 dimensi: Penerimaan Diri, Relasi Positif, Otonomi, Penguasaan Lingkungan, Tujuan Hidup, dan Pertumbuhan Pribadi.\n"
                  "• Bagian 2: 12 butir Identifikasi Diri / Ide Bunuh Diri (kategori: suicide_ideation) yang mengukur intensitas pikiran negatif dan risiko psikologis.")
    add_screenshot(doc, "admin_04_questions.png", "Gambar 1.5: Manajemen Bank Soal Asesmen Psikologis")

    add_heading_2(doc, "4.6 Pengaturan Skala Pilihan Jawaban (Option Groups)")
    add_paragraph(doc, "Menu 'Option Groups' (/admin/option-groups) mengatur pembobotan skala (Skala Likert 1–5: Sangat Tidak Sesuai s.d. Sangat Sesuai, atau Skala Frekuensi Tidak Pernah s.d. Selalu).")
    add_screenshot(doc, "admin_05_option_groups.png", "Gambar 1.6: Halaman Skala Pilihan Jawaban")

    add_heading_2(doc, "4.7 Manajemen Norma & Interpretasi Skor Asesmen")
    add_paragraph(doc, "Menu 'Interpretations' (/admin/interpretations) mengatur rentang skor, norma kategorisasi hasil asesmen, dan narasi rekomendasi tindakan klinis.")
    add_screenshot(doc, "admin_06_interpretations.png", "Gambar 1.7: Halaman Manajemen Interpretasi Skor")

    add_heading_2(doc, "4.8 Manajemen Modul Self-Help Budaya (Tribes & Materials)")
    add_paragraph(doc, "Menu 'Tribes' (/admin/tribes) memuat khazanah suku budaya Nusantara. Klik tombol 'Materials' untuk menyusun materi narasi inspiratif dan butir pertanyaan reflektif SFBC bagi siswa.")
    add_screenshot(doc, "admin_07_tribes.png", "Gambar 1.8: Daftar Suku Budaya Nusantara (Tribes)")
    add_screenshot(doc, "admin_08_materials.png", "Gambar 1.9: Manajemen Materi Refleksi SFBC pada Modul Terpilih")

    add_heading_2(doc, "4.9 Manajemen Modul Pembelajaran (Lessons)")
    add_paragraph(doc, "Menu 'Lessons' (/admin/lessons) digunakan untuk mempublikasikan artikel psikoedukasi, video motivasi, dan panduan keterampilan hidup yang dapat dipelajari secara mandiri oleh siswa.")
    add_screenshot(doc, "admin_09_lessons.png", "Gambar 1.10: Halaman Manajemen Modul Pembelajaran (Lessons)")

    output_path = os.path.join(DOCS_DIR, "Panduan_Pengguna_Ruang_Tumbuh_Admin.docx")
    doc.save(output_path)
    print(f"Berhasil membuat: {output_path}")

# ==============================================================================
# 2. GENERATE DOKUMEN KONSELOR
# ==============================================================================
def generate_konselor_doc():
    print("Membangun Dokumen Baru: Panduan_Pengguna_Ruang_Tumbuh_Konselor.docx...")
    doc = setup_document("Buku Panduan Konselor")
    add_cover(
        doc,
        title="Buku Panduan Pengguna: Konselor",
        subtitle="Platform Konseling Online, Skrining Resiliensi, & Peringatan Krisis Remaja",
        module_name="Modul Layanan Konseling Klinis & Pemantauan Binaan (Konselor)"
    )

    # 1. Judul & Versi
    add_heading_1(doc, "1. Informasi Modul & Akses Sistem")
    add_paragraph(doc, "Dokumen ini menyajikan panduan komprehensif bagi Konselor Ruang Tumbuh. Melalui pembaruan sistem terkini, modul Konselor dilengkapi fitur skrining klinis mutakhir berbasis instrumen Psychological Well-Being Carol D. Ryff (25 butir) dan Skala Deteksi Ide Bunuh Diri (12 butir), indikator peringatan krisis otomatis, manajemen sesi konseling daring (live chat & Google Meet/Zoom), serta lembar analisis komparasi dimensi psikologis.")
    
    headers_info = ["Komponen", "Spesifikasi / Keterangan"]
    data_info = [
        ["Nama Aplikasi", "Ruang Tumbuh"],
        ["Domain Akses", "https://ruangtumbuh.appocket.id"],
        ["Role Sasaran", "Konselor / Guru Bimbingan dan Konseling (BK)"],
        ["Modul Utama", "Peringatan Krisis, Skrining Prioritas, Live Chat, Telekonseling, & Evaluasi"],
        ["Periode Pembaruan", "Oktober 2026"],
    ]
    add_table_data(doc, headers_info, data_info, [Inches(2.0), Inches(4.4)], theme_color="0D9488")

    # 2. Pendahuluan Ringan
    add_heading_1(doc, "2. Pendahuluan Ringan")
    add_paragraph(doc, "Dalam platform Ruang Tumbuh, Konselor berperan sebagai garda terdepan dalam menjaga keselamatan dan menumbuhkan daya lenting (resiliensi) remaja. Sistem secara otomatis menganalisis jawaban asesmen siswa dan memetakan tingkat risiko agar konselor dapat segera memberikan respon yang tepat dan proporsional.")
    add_paragraph(doc, "Fitur-fitur unggulan modul Konselor terkini:\n"
                  "1. 4 Kartu Statistik Cepat: Memantau Total Konseli Terdaftar, Jadwal Sesi Aktif, Peringatan Krisis (Merah), dan Prioritas Skrining (Kuning <= 80).\n"
                  "2. Badge Peringatan Otomatis: Memberi tanda '⚠ Krisis' (merah mencolok) jika skor Ide Bunuh Diri >= 41 atau butir inti bernilai tinggi, serta badge 'Prioritas (<= 80)' jika kesejahteraan psikologis rendah.\n"
                  "3. Kolom Skor Terpisah: Menampilkan skor 'WB: {skor}/125' dan 'Ide: {skor}/60' secara transparan pada tabel daftar binaan.\n"
                  "4. Analisis Mendalam 6 Dimensi Well-Being & 4 Aspek Ide Bunuh Diri: Disertai norma empiris, titik potong skrining, dan protokol tindakan darurat klinis.\n"
                  "5. Ruang Sesi & Live Chat: Berbalas pesan langsung dan mengatur jadwal pertemuan tatap muka virtual (Google Meet / Zoom).\n"
                  "6. Laporan PDF Resmi: Mengunduh arsip asesmen siswa dalam format dokumen cetak resmi.")

    # 3. Prasyarat (Prerequisites)
    add_heading_1(doc, "3. Prasyarat (Prerequisites)")
    prereq_headers = ["Kebutuhan", "Spesifikasi Minimal yang Disarankan"]
    prereq_data = [
        ["Akun Konselor", "Akun aktif dengan role 'konselor' (didaftarkan oleh Administrator)."],
        ["Perangkat", "Laptop atau PC yang dilengkapi mikrofon dan kamera (webcam)."],
        ["Peramban Web", "Google Chrome atau Microsoft Edge versi terkini."],
        ["Koneksi Internet", "Minimal 5 Mbps stabil untuk mendukung telekonseling video tanpa hambatan."],
        ["Platform Eksternal", "Akun Google untuk penyelenggaraan Google Meet / Zoom."]
    ]
    add_table_data(doc, prereq_headers, prereq_data, [Inches(2.2), Inches(4.2)], theme_color="0D9488")

    # 4. Langkah Penggunaan (Step-by-Step)
    add_heading_1(doc, "4. Langkah Penggunaan (Step-by-Step)")
    
    add_heading_2(doc, "4.1 Masuk ke Halaman Login Konselor")
    add_paragraph(doc, "Kunjungi https://ruangtumbuh.appocket.id/login, masukkan email/username konselor dan kata sandi, lalu klik tombol 'Log in'.")
    add_screenshot(doc, "02_login_page.png", "Gambar 2.1: Halaman Masuk Akun Konselor")

    add_heading_2(doc, "4.2 Dashboard Konselor: Monitoring Skrining & Peringatan Krisis")
    add_paragraph(doc, "Dashboard Konselor (/konselor) kini dilengkapi 4 kartu statistik cerdas di bagian atas dan tabel daftar siswa binaan dengan indikator peringatan visual:")
    add_screenshot(doc, "konselor_01_dashboard.png", "Gambar 2.2: Tampilan Baru Dashboard Konselor dengan Peringatan Krisis & Skrining")
    
    headers_dashboard_guide = ["Elemen Antarmuka", "Keterangan Klinis & Fungsi"]
    data_dashboard_guide = [
        ["Kartu Peringatan Krisis (Merah)", "Jumlah siswa yang membutuhkan penanganan segera di hari yang sama (Skor Ide >= 41 atau butir inti frekuensi tinggi)."],
        ["Kartu Prioritas Skrining (Kuning)", "Jumlah siswa dengan skor Well-Being <= 80 (skala 125) yang masuk prioritas pencegahan selektif."],
        ["Kolom Skor Awal & Skor Akhir", "Menampilkan skor ganda: 'WB: {skor}/125' dan 'Ide: {skor}/60'."],
        ["Badge ⚠ Krisis & Prioritas", "Badge warna mencolok untuk memudahkan konselor menentukan urutan prioritas penanganan siswa."],
        ["Tombol 'Sesi & Chat'", "Membuka ruang komunikasi dan manajemen sesi konseling online."],
        ["Tombol 'Hasil Asesmen & PDF'", "Membuka lembar analisis lengkap 6 dimensi psikologis dan unduh laporan PDF."]
    ]
    add_table_data(doc, headers_dashboard_guide, data_dashboard_guide, [Inches(2.2), Inches(4.2)], theme_color="0D9488")

    add_heading_2(doc, "4.3 Ruang Sesi Konseling: Live Chat & Jadwal Virtual")
    add_paragraph(doc, "Klik tombol 'Sesi & Chat' untuk masuk ke ruang kerja konseling (/konselor/konseli/{id}/session). Header halaman menampilkan identitas siswa, status skor terbaru, serta banner darurat jika siswa terindikasi krisis.")
    add_screenshot(doc, "konselor_02_konseli_session.png", "Gambar 2.3: Ruang Sesi Konseling Konselor Lengkap dengan Chat & Penjadwalan Meet")
    add_callout(doc, "Lakukan penetapan jadwal temu dan sematkan link Google Meet setidaknya 1 jam sebelum sesi agar siswa menerima jadwal dengan jelas.", title="💡 TIPS PENJADWALAN KONSELI", border_color="0D9488", bg_color="F0FDFA")

    add_heading_2(doc, "4.4 Analisis Klinis Asesmen Psikologis & Laporan PDF Resmi")
    add_paragraph(doc, "Klik tombol 'Hasil Asesmen & PDF' (/konselor/konseli/{id}/wellbeing). Halaman ini menyajikan rincian komprehensif:\n"
                  "1. Ringkasan Skor & Titik Potong Skrining: Membandingkan perolehan skor siswa dengan norma populasi (M=91.70, SD=11.93).\n"
                  "2. Analisis 6 Dimensi Well-Being: Menampilkan skor pada dimensi Penerimaan Diri, Relasi Positif, Otonomi, Penguasaan Lingkungan, Tujuan Hidup, dan Pertumbuhan Pribadi.\n"
                  "3. Analisis 4 Aspek Ide Bunuh Diri: Membedah keputusasaan, perasaan menjadi beban, dan pikiran tentang kematian.\n"
                  "4. Protokol Tindakan Klinis: Panduan baku tindakan darurat, safety planning, dan rujukan profesional.\n"
                  "5. Unduh PDF & Reset: Tombol unduh laporan resmi untuk arsip sekolah serta reset asesmen jika diperlukan tes ulang.")
    add_screenshot(doc, "konselor_03_konseli_wellbeing.png", "Gambar 2.4: Lembar Analisis Klinis Asesmen Well-Being & Ide Bunuh Diri")
    add_callout(doc, "Jika siswa memiliki indikator '⚠ Krisis', segera lakukan kontak di hari yang sama, libatkan wali kelas/orang tua, dan susun rencana keselamatan (safety plan).", title="🚨 PROTOKOL TINDAKAN KRISIS", border_color="DC2626", bg_color="FEF2F2")

    add_heading_2(doc, "4.5 Memeriksa Jawaban Refleksi Budaya Konseli")
    add_paragraph(doc, "Melalui tombol 'Jawaban Self-Help' (/konselor/konseli/{id}/self-help), konselor dapat meninjau narasi reflektif siswa terhadap modul kearifan lokal yang telah dikerjakannya sebagai bahan eksplorasi solusi singkat (SFBC).")
    add_screenshot(doc, "konselor_04_konseli_selfhelp.png", "Gambar 2.5: Pemantauan Lembar Refleksi Mandiri Budaya Siswa")

    output_path = os.path.join(DOCS_DIR, "Panduan_Pengguna_Ruang_Tumbuh_Konselor.docx")
    doc.save(output_path)
    print(f"Berhasil membuat: {output_path}")

# ==============================================================================
# 3. GENERATE DOKUMEN KONSELI (SISWA)
# ==============================================================================
def generate_konseli_doc():
    print("Membangun Dokumen Baru: Panduan_Pengguna_Ruang_Tumbuh_Konseli.docx...")
    doc = setup_document("Buku Panduan Siswa (Konseli)")
    add_cover(
        doc,
        title="Buku Panduan Pengguna: Siswa (Konseli)",
        subtitle="Ruang Aman Remaja Bercerita, Bertumbuh, & Membangun Resiliensi Diri",
        module_name="Modul Layanan Konseling Mandiri & Evaluasi Perkembangan (Konseli)"
    )

    # 1. Judul & Versi
    add_heading_1(doc, "1. Informasi Modul & Akses Sistem")
    add_paragraph(doc, "Buku panduan ini disusun khusus untuk Siswa / Remaja (Konseli) agar dapat memanfaatkan platform Ruang Tumbuh dengan nyaman, aman, dan mandiri. Panduan ini mencakup langkah pendaftaran akun, pengisian asesmen resiliensi diri (Well-Being & Deteksi Pikiran Negatif), pemilihan konselor pendamping, percakapan konseling daring, hingga melihat perkembangan kemajuan diri melalui halaman Evaluasi & Tindak Lanjut.")
    
    headers_info = ["Komponen", "Spesifikasi / Keterangan"]
    data_info = [
        ["Nama Aplikasi", "Ruang Tumbuh"],
        ["Domain Akses", "https://ruangtumbuh.appocket.id"],
        ["Role Sasaran", "Siswa / Remaja (Konseli)"],
        ["Fitur Unggulan", "Asesmen Diri, Pilih Konselor, Chat & Video Meet, Grafik Evaluasi, & Modul Budaya"],
        ["Periode Pembaruan", "Oktober 2026"],
    ]
    add_table_data(doc, headers_info, data_info, [Inches(2.0), Inches(4.4)], theme_color="2563EB")

    # 2. Pendahuluan Ringan
    add_heading_1(doc, "2. Pendahuluan Ringan")
    add_paragraph(doc, "Selamat datang di Ruang Tumbuh! Aplikasi ini dirancang sebagai tempat curhat yang aman, nyaman, dan bebas stigma untuk membantumu mengatasi stres belajar, kecemasan, rasa tertekan, atau masalah pertemanan.")
    add_paragraph(doc, "Keuntungan menggunakan Ruang Tumbuh:\n"
                  "1. Mengetahui Kondisi Kesejahteraan Mentalmu: Melalui asesmen terstandar (skala 125 poin), kamu bisa mengetahui kekuatan diri dan hal-hal yang perlu kamu tingkatkan.\n"
                  "2. Konselor Pendamping yang Ramah: Kamu bebas memilih guru BK atau konselor profesional yang paling membuatmu nyaman.\n"
                  "3. Percakapan Tertutup & Rahasia: Seluruh curhatanmu dijamin kerahasiaannya sesuai kode etik profesi konseling.\n"
                  "4. Grafik Evaluasi Kemajuan: Kamu bisa memantau peningkatan kesejahteraan psikologismu sebelum dan sesudah mengikuti bimbingan konseling.\n"
                  "5. Modul Kearifan Budaya Nusantara: Cerita inspiratif dan petuah bijak dari suku nusantara untuk melatih ketangguhan mentalmu.")
    add_callout(doc, "Jangan ragu untuk menceritakan kondisi yang sebenarnya. Konselor hadir untuk mendengarkan, memahami, dan membantumu menemukan solusi terbaik tanpa menghakimi.", title="🛡️ PRIVASI & KERAHASIAAN AMAN", border_color="2563EB", bg_color="EFF6FF")

    # 3. Prasyarat (Prerequisites)
    add_heading_1(doc, "3. Prasyarat (Prerequisites)")
    prereq_headers = ["Kebutuhan", "Penjelasan"]
    prereq_data = [
        ["Akun Siswa", "Daftar mandiri melalui formulir registrasi online."],
        ["Perangkat", "Smartphone (Android / iPhone), Tablet, atau Laptop/PC."],
        ["Peramban (Browser)", "Google Chrome, Safari, atau Microsoft Edge."],
        ["Koneksi Internet", "Akses internet seluler atau Wi-Fi yang stabil."],
        ["Aplikasi Pendukung", "Aplikasi Google Meet jika ingin melakukan panggilan video tatap muka lewat HP."]
    ]
    add_table_data(doc, prereq_headers, prereq_data, [Inches(2.2), Inches(4.2)], theme_color="2563EB")

    # 4. Langkah Penggunaan (Step-by-Step)
    add_heading_1(doc, "4. Langkah Penggunaan (Step-by-Step)")
    
    add_heading_2(doc, "4.1 Pendaftaran Akun Siswa Baru (Registrasi)")
    add_paragraph(doc, "1. Buka tautan: https://ruangtumbuh.appocket.id/register\n"
                  "2. Masukkan Nama Lengkap, Username, Alamat Email aktif, Asal Sekolah, Kelas, dan Jenis Kelamin.\n"
                  "3. Buat kata sandi akunmu, lalu klik 'Daftar Sekarang'.")
    add_screenshot(doc, "03_register_page.png", "Gambar 3.1: Halaman Registrasi Akun Siswa Baru")

    add_heading_2(doc, "4.2 Masuk ke Akun (Login)")
    add_paragraph(doc, "Kunjungi https://ruangtumbuh.appocket.id/login, masukkan email/username dan kata sandi, lalu klik tombol 'Masuk'.")
    add_screenshot(doc, "02_login_page.png", "Gambar 3.2: Halaman Masuk Akun Siswa")

    add_heading_2(doc, "4.3 Mengenal Dashboard Siswa")
    add_paragraph(doc, "Setelah masuk, Dashboard Siswa (/konseli) menyajikan panduan langkah bertumbuh (Langkah 1 s.d. Langkah 7). Terdapat kartu pemantauan asesmen awal, tombol pintas unduh PDF resmi, serta akses ke seluruh layanan konseling.")
    add_screenshot(doc, "konseli_01_dashboard.png", "Gambar 3.3: Dashboard Utama Siswa Ruang Tumbuh Terkini")

    add_heading_2(doc, "4.4 Mengisi Asesmen Resiliensi Diri (Well-Being)")
    add_paragraph(doc, "Masuk ke menu 'Asesmen' (/konseli/wellbeing). Pengisian asesmen terbagi menjadi dua bagian dengan panduan psikologis yang menenangkan:\n"
                  "• Bagian 1: Pernyataan mengenai kesejahteraan psikologis sehari-hari (25 butir).\n"
                  "• Bagian 2: Pernyataan mengenai perasaan mendalam dan pikiran negatif (12 butir).\n"
                  "Isi setiap pernyataan dengan jujur. Setelah selesai, kamu bisa langsung melihat hasil dan mengunduh berkas laporan PDF.")
    add_screenshot(doc, "konseli_02_wellbeing_index.png", "Gambar 3.4: Menu Asesmen Resiliensi Siswa")

    add_heading_2(doc, "4.5 Memilih Konselor Pendamping")
    add_paragraph(doc, "Pada menu 'Pilih Konselor' (/konseli/counselor), kamu dapat melihat profil para konselor, riwayat pendidikan, motto hidup, serta nomor kontak. Klik tombol 'Pilih Konselor Ini' pada profil konselor yang kamu kehendaki.")
    add_screenshot(doc, "konseli_03_counselor_select.png", "Gambar 3.5: Halaman Pemilihan Konselor Pendamping")

    add_heading_2(doc, "4.6 Ruang Konseling Online: Live Chat & Jadwal Video Meet")
    add_paragraph(doc, "Setelah memiliki konselor, buka menu 'Konseling Online' (/konseli/counseling):\n"
                  "1. Ketik pesan curhatmu di kotak live chat untuk memulai obrolan.\n"
                  "2. Konselor akan membalas dan menetapkan jadwal tatap muka daring.\n"
                  "3. Saat jadwal tiba, klik tombol 'Masuk ke Video Meet' untuk terhubung langsung melalui Google Meet.")
    add_screenshot(doc, "konseli_04_counseling_room.png", "Gambar 3.6: Ruang Percakapan Chat & Jadwal Pertemuan Daring")

    add_heading_2(doc, "4.7 Mengisi Lembar Refleksi Pasca Konseling")
    add_paragraph(doc, "Setelah sesi berbincang selesai, buka menu 'Refleksi Sesi' (/konseli/reflection) untuk menuliskan hal-hal positif yang kamu pelajari dan komitmen aksi perbaikan dirimu.")
    add_screenshot(doc, "konseli_05_reflection.png", "Gambar 3.7: Lembar Refleksi Pasca Sesi Konseling")

    add_heading_2(doc, "4.8 Halaman Evaluasi & Tindak Lanjut: Memantau Kemajuan Diri")
    add_paragraph(doc, "Menu 'Evaluasi & Tindak Lanjut' (/konseli/evaluation) menampilkan grafik visual perbandingan hasil asesmen awal (pre-test) dan asesmen sesudah (post-test):\n"
                  "• Grafik Komparasi Well-Being (skala 125) dan Penurunan Risiko Pikiran Negatif (skala 60).\n"
                  "• Banner Indikator Evaluasi: Memberikan umpan balik apakah risiko menurun dan resiliensimu meningkat secara positif.\n"
                  "• Tabel Rincian 6 Dimensi Well-Being: Menampilkan delta perolehan poin per dimensi (Penerimaan Diri, Relasi Positif, Otonomi, Penguasaan Lingkungan, Tujuan Hidup, Pertumbuhan Pribadi).\n"
                  "• Catatan & Rekomendasi Konselor: Membaca saran perbaikan yang dituliskan langsung oleh konselormu.")
    add_screenshot(doc, "konseli_06_evaluation.png", "Gambar 3.8: Tampilan Baru Grafik Evaluasi & Komparasi Dimensi Psikologis")

    add_heading_2(doc, "4.9 Modul Self-Help Kearifan Budaya Nusantara")
    add_paragraph(doc, "Ruang Tumbuh menyediakan petuah bijak dari berbagai suku Nusantara (/konseli/tribe). Pelajari kisah ketangguhan leluhur dan jawab pertanyaan reflektif mandiri untuk menguatkan daya lentingmu.")
    add_screenshot(doc, "konseli_07_tribe_select.png", "Gambar 3.9: Halaman Pilihan Suku Budaya Nusantara")
    add_screenshot(doc, "konseli_08_selfhelp_detail.png", "Gambar 3.10: Materi Cerita & Lembar Refleksi Mandiri Budaya")

    add_heading_2(doc, "4.10 Modul Pembelajaran & Pengaturan Profil")
    add_paragraph(doc, "Kamu dapat membaca materi tips kesehatan mental di menu 'Pembelajaran' (/konseli/lessons) serta memperbarui data diri di menu 'Profil' (/konseli/profile).")
    add_screenshot(doc, "konseli_09_lessons.png", "Gambar 3.11: Modul Pembelajaran Edukatif")
    add_screenshot(doc, "konseli_10_profile.png", "Gambar 3.12: Pengaturan Profil Siswa")

    output_path = os.path.join(DOCS_DIR, "Panduan_Pengguna_Ruang_Tumbuh_Konseli.docx")
    doc.save(output_path)
    print(f"Berhasil membuat: {output_path}")

# ==============================================================================
# MAIN EXECUTION
# ==============================================================================
if __name__ == "__main__":
    generate_admin_doc()
    generate_konselor_doc()
    generate_konseli_doc()
    print("\nSeluruh dokumen Word (.docx) versi terbaru berhasil dibangun dengan sempurna!")
